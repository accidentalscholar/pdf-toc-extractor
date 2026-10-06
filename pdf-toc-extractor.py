# ==========================================
# PDF TOC Extractor
# Version: 1.3.0
# Citation: Pundir, V. (2026, October 06). PDF TOC Extractor Version (1.3.0). Retrieved from https://github.com/accidentalscholar/pdf-toc-extractor. 
# Citation: RIS and BibTeX files included for referencing software.
# Tested in: Python 3.10.9 64 bit packaged by Anaconda, Inc.
# Reporsitory: https://github.com/accidentalscholar/pdf-toc-extractor
# Provided under: GNU AFFERO GENERAL PUBLIC LICENSE (see accompanying license file)
# ==========================================

import sys
import os
import subprocess
import gc
import re

def setup_environment():
    """
    Checks for a path.txt file to add custom paths to sys.path, 
    and auto-installs missing required libraries.
    """
    # 1. Check for path.txt in the current working directory
    if os.path.exists("path.txt"):
        with open("path.txt", "r") as f:
            for line in f:
                custom_path = line.strip()
                if custom_path and os.path.exists(custom_path):
                    sys.path.append(custom_path)
                    print(f"Added custom path: {custom_path}")

    # 2. Auto-install required packages
    required_packages = {
        'pymupdf': 'fitz',
        'openpyxl': 'openpyxl'
    }
    
    for pkg, module in required_packages.items():
        try:
            __import__(module)
        except ImportError:
            print(f"Library '{module}' not found. Installing '{pkg}'...")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", pkg])
                print(f"Successfully installed {pkg}.")
            except Exception as e:
                print(f"Failed to install {pkg}. Error: {e}")
                sys.exit(1)

# Run environment setup before importing the third-party libraries
setup_environment()

import fitz  # PyMuPDF
import tkinter as tk
from tkinter import filedialog
from openpyxl import Workbook, load_workbook

# List of common front/back matter keywords to exclude from final output
EXCLUDE_KEYWORDS = [
    'index', 'appendix', 'bibliography', 'references', 'front matter', 
    'back matter', 'acknowledgments', 'preface', 'about the author', 
    'glossary', 'epilogue', 'prologue', 'about the contributors',
    'title page', 'copyright', 'contents', 'table of contents',
    'discussion guide', 'notes', 'front cover', 'copyright page',
    'back cover'
]

def clean_entries(entries):
    """Filters out unwanted front/back matter and index entries."""
    cleaned = []
    for title, page in entries:
        if not any(kw in str(title).lower() for kw in EXCLUDE_KEYWORDS):
            cleaned.append((title, page))
    return cleaned

def extract_toc_from_text(doc, toc_bookmarks):
    """
    Attempts to find the TOC pages based on bookmarks and extracts entries 
    by reading the physical text on those pages.
    """
    toc_start_page = -1
    toc_end_page = -1
    
    # Locate the TOC bookmark to find the physical pages
    for i, item in enumerate(toc_bookmarks):
        lvl, title, page = item
        if re.search(r'(?i)^(contents|table of contents|toc)$', title.strip()):
            toc_start_page = page
            # Find the very next bookmark at the same or higher hierarchy to determine end of TOC
            for j in range(i + 1, len(toc_bookmarks)):
                if toc_bookmarks[j][0] <= lvl:
                    toc_end_page = toc_bookmarks[j][2]
                    break
            break
            
    if toc_start_page == -1:
        return [] # Could not locate physical TOC pages via bookmarks
        
    if toc_end_page == -1:
        # If no subsequent bookmark, just scan the next 3 pages
        toc_end_page = toc_start_page + 3
        
    # PyMuPDF uses 0-based page indexing, get_toc() returns 1-based
    start_idx = max(0, toc_start_page - 1)
    end_idx = min(toc_end_page - 1, start_idx + 5, doc.page_count - 1)
    if end_idx < start_idx:
        end_idx = start_idx + 2
        
    extracted_text_entries = []
    
    for p_idx in range(start_idx, end_idx + 1):
        if p_idx >= doc.page_count: break
        page = doc[p_idx]
        text = page.get_text("text")
        
        for line in text.split('\n'):
            line = line.strip()
            if not line:
                continue
            
            # Match text followed by spaces or dots, ending with digits (e.g., "Chapter 1 ...... 12")
            match = re.search(r'^(.+?)(?:\.{2,}|\s{2,}|\s+)(\d+)$', line)
            if match:
                ch_title = match.group(1).strip(' ._-')
                ch_page = match.group(2)
                
                # Exclude lines that are just numbers or very short noisy fragments
                if len(ch_title) > 2 and not ch_title.isdigit():
                    extracted_text_entries.append((ch_title, ch_page))
                    
    return extracted_text_entries

def extract_toc_from_bookmarks(toc_bookmarks):
    """
    Applies the specific fallback logic based on page gaps between levels
    to avoid extracting both 'Parts' and 'Chapters'.
    """
    level1_items = []
    level2_items = []
    parents_with_first_child = []
    
    current_l1 = None
    for item in toc_bookmarks:
        lvl, title, page = item
        if lvl == 1:
            current_l1 = item
            level1_items.append(item)
        elif lvl == 2:
            level2_items.append(item)
            # Track the first level 2 child for the current level 1 parent
            if current_l1 is not None and not any(p == current_l1 for p, c in parents_with_first_child):
                parents_with_first_child.append((current_l1, item))
                
    if not level1_items:
        return [(item[1], item[2]) for item in toc_bookmarks]
        
    if not parents_with_first_child:
        return [(item[1], item[2]) for item in level1_items]
        
    # User Logic: If page gap between Top Level and its first Second Level is ALWAYS < 3
    # Top Level = Parts (Ignore). Second Level = Chapters (Use).
    all_less_than_3 = True
    for p, c in parents_with_first_child:
        parent_page = p[2]
        child_page = c[2]
        if parent_page > 0 and child_page > 0:
            if abs(child_page - parent_page) >= 3:
                all_less_than_3 = False
                break
                
    if all_less_than_3:
        final_toc = level2_items
    else:
        final_toc = level1_items
        
    return [(item[1], item[2]) for item in final_toc]

def process_pdfs():
    """
    Main function to select directory, orchestrate extraction attempts, 
    and save outputs progressively to Excel.
    """
    root = tk.Tk()
    root.withdraw()
    
    folder_path = filedialog.askdirectory(title="Select Folder Containing PDFs")
    
    # Safely close the tkinter window for Spyder compatibility
    root.update()
    root.destroy()
    
    if not folder_path:
        print("No folder was selected. Exiting script.")
        return

    excel_filename = "Extracted_PDF_TOCs.xlsx"
    excel_path = os.path.join(folder_path, excel_filename)

    if not os.path.exists(excel_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "TOC Data"
        ws.append(["Name of book", "Name of chapter/article", "Page number", "Filename"])
        wb.save(excel_path)
        wb.close()

    pdf_files = [f for f in os.listdir(folder_path) if f.lower().endswith('.pdf')]
    
    if not pdf_files:
        print(f"No PDF files found in the directory: {folder_path}")
        return

    print(f"Found {len(pdf_files)} PDF(s). Starting extraction...\n")

    for filename in pdf_files:
        pdf_path = os.path.join(folder_path, filename)
        print(f"Processing: {filename}")
        
        doc = None
        wb = None
        
        try:
            doc = fitz.open(pdf_path)
            
            book_name = doc.metadata.get("title")
            if not book_name or book_name.strip() == "":
                book_name = os.path.splitext(filename)[0]
            
            toc_bookmarks = doc.get_toc()
            final_entries = []
            
            # Attempt 1: Primary extraction using Intelligent Bookmark Logic
            if toc_bookmarks:
                bookmark_entries = extract_toc_from_bookmarks(toc_bookmarks)
                final_entries = clean_entries(bookmark_entries)
                if final_entries:
                    print("  -> Successfully extracted TOC using intelligent embedded bookmark logic.")
            
            # Attempt 2: Fallback to Physical Text Extraction if bookmarks yield no results
            if not final_entries:
                text_entries = extract_toc_from_text(doc, toc_bookmarks)
                final_entries = clean_entries(text_entries)
                if final_entries:
                    print("  -> Embedded bookmarks yielded no results. Used physical text fallback.")
                
            wb = load_workbook(excel_path)
            ws = wb.active
            
            if final_entries:
                for title, page in final_entries:
                    ws.append([book_name, title, page, filename])
                print(f"  -> Extracted {len(final_entries)} chapter/article entries.")
            else:
                ws.append([book_name, "NO VALID TOC FOUND", "N/A", filename])
                print("  -> No embedded bookmarks or valid text TOC found.")
            
            wb.save(excel_path)
            
        except Exception as e:
            print(f"  -> Error processing {filename}: {e}")
            
        finally:
            if wb is not None:
                wb.close()
            if doc is not None:
                doc.close()
            
            # Flush memory progressively
            gc.collect()

    print(f"\nProcessing complete! Excel file saved progressively to:\n{excel_path}")

if __name__ == "__main__":
    print(f"PDF TOC Extractor v{__version__}")
    process_pdfs()